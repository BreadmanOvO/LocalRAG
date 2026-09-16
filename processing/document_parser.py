"""Format-specific extraction shared by the asset API and ingestion scripts."""
from __future__ import annotations

import csv
import io
import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Callable


class DocumentParseError(ValueError):
    pass


@dataclass(frozen=True)
class DocumentPart:
    text: str
    locator: str
    metadata: dict


SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".json", ".pdf", ".docx", ".xlsx", ".png", ".jpg", ".jpeg", ".webp"}


def decode_text(content: bytes) -> str:
    encodings = ["utf-16"] if content.startswith((b"\xff\xfe", b"\xfe\xff")) else ["utf-8-sig", "gb18030"]
    for encoding in encodings:
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("无法识别文本编码，请转换为 UTF-8 后上传")


def parse_document(filename: str, content: bytes, *, vision: Callable[[bytes, str], str] | None = None) -> list[DocumentPart]:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        raise DocumentParseError("不支持此格式；旧版 Word/Excel 请先另存为 DOCX/XLSX")
    if not content:
        raise DocumentParseError("文件为空")
    if suffix in {".docx", ".xlsx"}:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(item.file_size for item in archive.infolist()) > 200 * 1024 * 1024:
                raise DocumentParseError("文档解压后超过 200 MB，请拆分后上传")
    parts: list[DocumentPart] = []
    if suffix in {".txt", ".md"}:
        parts.append(DocumentPart(decode_text(content), "全文", {"format": suffix[1:]}))
    elif suffix in {".csv", ".tsv"}:
        rows = csv.reader(io.StringIO(decode_text(content)), delimiter="\t" if suffix == ".tsv" else ",")
        headers = next(rows, [])
        for number, row in enumerate(rows, 2):
            if any(cell.strip() for cell in row):
                text = "\n".join(f"{headers[i] if i < len(headers) else f'列{i+1}'}: {value}" for i, value in enumerate(row))
                parts.append(DocumentPart(text, f"第 {number} 行", {"row": number, "format": "table"}))
    elif suffix == ".json":
        data = json.loads(decode_text(content))
        values = list(enumerate(data)) if isinstance(data, list) else [("$", data)]
        parts.extend(DocumentPart(json.dumps(value, ensure_ascii=False, indent=2), f"JSON {key}", {"json_path": str(key), "format": "json"}) for key, value in values)
    elif suffix == ".pdf":
        import fitz
        with fitz.open(stream=content, filetype="pdf") as document:
            if document.needs_pass:
                raise DocumentParseError("PDF 已加密，请上传解密后的文件")
            for number, page in enumerate(document, 1):
                text = page.get_text(sort=True).strip()
                if not text:
                    if not page.get_images() and not page.get_drawings():
                        continue
                    if vision is None:
                        raise DocumentParseError(f"第 {number} 页没有可提取文本，需要启用支持图片的模型后重试")
                    text = vision(page.get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).tobytes("png"), "image/png")
                parts.append(DocumentPart(text, f"第 {number} 页", {"page_start": number, "page_end": number, "format": "pdf"}))
    elif suffix == ".docx":
        from docx import Document
        document = Document(io.BytesIO(content))
        for number, paragraph in enumerate(document.paragraphs, 1):
            if paragraph.text.strip():
                parts.append(DocumentPart(paragraph.text, f"段落 {number}", {"paragraph": number, "format": "docx"}))
        for index, table in enumerate(document.tables, 1):
            headers = [cell.text for cell in table.rows[0].cells]
            for number, row in enumerate(table.rows[1:], 2):
                text = "\n".join(f"{headers[i] if i < len(headers) else f'列{i+1}'}: {cell.text}" for i, cell in enumerate(row.cells))
                parts.append(DocumentPart(text, f"表 {index} 行 {number}", {"table": index, "row": number, "format": "table"}))
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        try:
            for sheet in workbook:
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, ())
                for number, row in enumerate(rows, 2):
                    if any(value is not None for value in row):
                        text = "\n".join(f"{(headers[i] if i < len(headers) else None) or f'列{i+1}'}: {value if value is not None else ''}" for i, value in enumerate(row))
                        parts.append(DocumentPart(text, f"{sheet.title} 第 {number} 行", {"sheet": sheet.title, "row": number, "format": "table"}))
        finally:
            workbook.close()
    else:
        if vision is None:
            raise DocumentParseError("图片解析需要启用支持图片的模型，请在模型设置中配置后重试")
        media = "image/jpeg" if suffix in {".jpg", ".jpeg"} else f"image/{suffix[1:]}"
        parts.append(DocumentPart(vision(content, media), "图片", {"format": "image", "extraction": "vlm"}))
    parts = [part for part in parts if part.text.strip()]
    if not parts:
        raise DocumentParseError("文件中没有可索引的正文或数据行")
    return parts
