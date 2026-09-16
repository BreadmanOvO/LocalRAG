"""Format boundaries, persistent jobs, index rollback and API space isolation."""
import base64
import io
import json
from pathlib import Path
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient

from agent_platform.api import create_app
from agent_platform.capability_packs import LocalObjectStore
from agent_platform.capability_packs.asset_ingestion import AssetIngestionService
from core.chunking import chunk_text_baseline
from core.ingestion_workflow import IngestionWorkflow
from processing.document_parser import DocumentParseError, parse_document


@pytest.mark.parametrize("filename,content,locator", [
    ("notes.txt", "中文正文".encode("gb18030"), "全文"),
    ("rows.csv", "姓名,部门\n小明,研发".encode(), "第 2 行"),
    ("rows.tsv", b"name\tdepartment\nalice\tengineering", "第 2 行"),
    ("rows.json", b'[{"name":"alice"}]', "JSON 0"),
])
def test_text_formats(filename, content, locator):
    parts = parse_document(filename, content)
    assert len(parts) == 1 and parts[0].locator == locator
    assert "�" not in parts[0].text


def test_office_and_pdf_locators():
    from docx import Document
    from openpyxl import Workbook
    import fitz
    document = Document()
    document.add_paragraph("项目说明")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("名称", "数量"), ("相机", "6")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    buffer = io.BytesIO()
    document.save(buffer)
    parts = parse_document("a.docx", buffer.getvalue())
    assert parts[0].metadata["paragraph"] == 1
    assert "数量: 6" in parts[1].text
    workbook = Workbook()
    workbook.active.append(["名称", "数量"])
    workbook.active.append(["相机", 6])
    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    assert "数量: 6" in parse_document("a.xlsx", buffer.getvalue())[0].text
    with fitz.open() as pdf:
        pdf.new_page().insert_text((72, 72), "BEVFormer uses six cameras")
        pdf.new_page()  # A blank page should not require a vision model.
        parts = parse_document("a.pdf", pdf.tobytes())
    assert len(parts) == 1 and parts[0].metadata["page_start"] == 1


def test_image_requires_vision_and_invalid_files_fail():
    with pytest.raises(DocumentParseError, match="图片"):
        parse_document("image.png", b"image")
    parts = parse_document("image.png", b"image", vision=lambda content, media: "相机图示")
    assert parts[0].metadata["extraction"] == "vlm"
    for name, content in [("old.doc", b"old"), ("blank.txt", b""), ("blank.csv", b"title\n")]:
        with pytest.raises(DocumentParseError):
            parse_document(name, content)


class MemoryIndex:
    def __init__(self, root):
        self.rows = {}
        self._persist_directory = str(root)
        self._collection = SimpleNamespace(name="assets")
        self.fail_once = False

    def get(self, ids):
        return {"ids": [key for key in ids if key in self.rows]}

    def add_texts(self, *, texts, metadatas, ids):
        for key, text, metadata in zip(ids, texts, metadatas):
            self.rows[key] = (text, metadata)
            if self.fail_once:
                self.fail_once = False
                raise RuntimeError("provider secret must not be exposed")

    def delete(self, *, ids):
        for key in ids:
            self.rows.pop(key, None)

    def similarity_search(self, query, k):
        return [SimpleNamespace(page_content=text, metadata=metadata) for text, metadata in list(self.rows.values())[:k]]


def pipeline(tmp_path):
    workflows = {}
    def factory(space):
        if space not in workflows:
            root = tmp_path / space
            index = MemoryIndex(root)
            kb = SimpleNamespace(chroma=index, _chunk_upload=lambda text, metadata, **kw: chunk_text_baseline(text, source_metadata=metadata), add_chunk_records=lambda chunks, ids: index.add_texts(texts=[c.text for c in chunks], metadatas=[c.metadata for c in chunks], ids=ids))
            workflows[space] = IngestionWorkflow(knowledge_base=kb, staging_directory=root / "staging", registry_path=root / "sources.json", active_profile_path=root / "profile.json", uploaded_documents_directory=root / "docs", manifest_builder=lambda **kw: {"registry_source_count": 1, "chunk_count": len(index.rows), "corpus_fingerprint": "a" * 64, "registry_fingerprint": "b" * 64})
        return workflows[space]
    store = LocalObjectStore(tmp_path / "objects")
    return AssetIngestionService(store, root=tmp_path / "jobs", workflow_factory=factory), factory


def finish(service, job):
    service._futures[job["job_id"]].result(timeout=10)
    return service.get(job["job_id"])


def test_publish_retry_partial_rollback_and_persistence(tmp_path):
    service, factory = pipeline(tmp_path)
    factory("space-a").knowledge_base.chroma.fail_once = True
    job = finish(service, service.submit("x.csv", b"name,value\nalice,6", "space-a"))
    assert job["status"] == "failed" and job["stage"] == "indexing"
    assert "secret" not in job["error"]
    assert factory("space-a").knowledge_base.chroma.rows == {}
    job = finish(service, service.retry(job["job_id"]))
    assert job["status"] == "completed" and job["chunk_count"] > 0
    hits = service.search("space-a", "alice")
    assert hits and hits[0]["metadata"]["row"] == 2
    assert service.search("space-b", "alice") == []
    duplicate = finish(service, service.submit("copy.csv", b"name,value\nalice,6", "space-a"))
    assert duplicate["source_id"] == job["source_id"]
    assert len(factory("space-a").knowledge_base.chroma.rows) == job["chunk_count"]
    service.close()
    restored = AssetIngestionService(service.object_store, root=service.root, workflow_factory=factory)
    assert len(restored.list("space-a")) == 2
    assert restored.get(job["job_id"])["history"] == job["history"]
    restored.close()


def test_api_auth_and_real_pipeline_state(tmp_path):
    service, _ = pipeline(tmp_path)
    app = create_app(asset_store=service.object_store, auth_required=True, auth_tokens={"a": "space-a", "b": "space-b"})
    app.state.asset_ingestion = service
    a, b = {"Authorization": "Bearer a"}, {"Authorization": "Bearer b"}
    with TestClient(app) as client:
        body = {"space_id": "space-a", "filename": "hello.txt", "content_base64": base64.b64encode(b"hello searchable source").decode()}
        assert client.post("/asset-ingestions", json=body, headers=b).status_code == 403
        response = client.post("/asset-ingestions", json=body, headers=a)
        assert response.status_code == 202, response.text
        job = finish(service, response.json())
        assert job["status"] == "completed"
        url = f"/asset-ingestions/{job['job_id']}"
        assert client.get(url, headers=b).status_code == 403
        assert client.post(url + "/retry", headers=b).status_code == 403
        assert client.get("/assets/search", params={"space_id": "space-a", "q": "hello"}, headers=b).status_code == 403
        assert client.get("/assets/search", params={"space_id": "space-a", "q": "hello"}, headers=a).json()["items"]
        assert client.get("/asset-ingestions", params={"space_id": "space-a"}, headers=a).json()["items"][0]["status"] == "completed"


def test_restart_recovers_queued_job(tmp_path):
    service, factory = pipeline(tmp_path)
    job = finish(service, service.submit("a.txt", b"recover me", "space-a"))
    service.close()
    job.update(status="running", stage="indexing")
    service._save(job)
    restored = AssetIngestionService(service.object_store, root=service.root, workflow_factory=factory)
    assert finish(restored, job)["status"] == "completed"
    restored.close()
